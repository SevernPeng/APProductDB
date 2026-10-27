from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from catalog.models import Brand, Category, Product, ProductSpec, SpecDefinition
from comparison.services import build_comparison_rows, load_comparison_products


class WorkbookComparisonTests(TestCase):
    def setUp(self):
        brand = Brand.objects.create(name="TP-Link", slug="tp-link", is_own_brand=True)
        self.category = Category.objects.create(name="Managed Switch", slug="managed-switches")
        self.products = [Product.objects.create(brand=brand, category=self.category,
                         model=model, official_url=f"https://example.com/{model}")
                         for model in ("SG3428X", "SG3452X")]
        self.capacity = self.definition("switching_capacity_gbps", "Switching Capacity", "decimal", "Gbps")
        self.interfaces = self.definition("ethernet_interfaces", "Ethernet Interfaces")
        self.console = self.definition("console_interfaces", "Console Ports")
        self.buffer = self.definition("packet_buffer_mb", "Packet Buffer", "decimal", "MB")
        self.definition("poe_budget_w", "Total PoE Budget", "decimal", "W")
        self.matrix(self.products[0], "matrix-access-016", "Switching Capacity", "128 Gbps")
        self.matrix(self.products[1], "matrix-campus-021", "Switching Capacity", "176 Gbps")
        self.matrix(self.products[0], "matrix-access-005", "Interface", "24 GE RJ45 Ports\n4 SFP+ Slots\n1 USB-C Console Port")
        self.matrix(self.products[0], "matrix-access-014", "Packet Buffer Memory", "12 Mbit")
        for product in self.products:
            for definition in (self.capacity, self.interfaces, self.console, self.buffer):
                ProductSpec.objects.create(product=product, definition=definition, value_status="unknown")

    def definition(self, code, name, data_type="text", unit=""):
        return SpecDefinition.objects.create(code=code, display_name=name,
            data_type=data_type, unit=unit, group="Performance", category=self.category)

    def matrix(self, product, code, label, value):
        definition = self.definition(code, label)
        definition.is_core = False
        definition.save()
        return ProductSpec.objects.create(product=product, definition=definition,
            value_text=value, source_url=f"https://example.com/source/{product.pk}")

    def rows(self, **options):
        products = load_comparison_products([p.pk for p in self.products])
        return {row["code"]: row for row in build_comparison_rows(products, **options)}

    def test_unknown_placeholders_fall_back_without_writes_or_cross_product_leak(self):
        before = list(ProductSpec.objects.order_by("pk").values())
        rows = self.rows()
        self.assertEqual([v["display"] for v in rows["switching_capacity_gbps"]["values"]],
                         ["128 Gbps", "176 Gbps"])
        self.assertEqual(rows["console_interfaces"]["values"][0]["display"], "1 USB-C Console Port")
        self.assertEqual(rows["console_interfaces"]["values"][1]["display"], "Unknown")
        self.assertEqual(rows["packet_buffer_mb"]["values"][0]["display"], "12 Mbit")
        self.assertEqual(rows["switching_capacity_gbps"]["values"][0]["source_url"],
                         f"https://example.com/source/{self.products[0].pk}")
        self.assertEqual(list(ProductSpec.objects.order_by("pk").values()), before)

    def test_published_zero_and_not_applicable_are_preserved(self):
        ProductSpec.objects.filter(product=self.products[0], definition=self.capacity).update(
            value_status="published", value_number=Decimal("0"))
        ProductSpec.objects.filter(product=self.products[1], definition=self.capacity).update(
            value_status="not_applicable")
        row = self.rows()["switching_capacity_gbps"]
        self.assertEqual([v["display"] for v in row["values"]], ["0 Gbps", "Not Applicable"])

    def test_workbook_only_rows_outside_template_are_hidden(self):
        for i, product in enumerate(self.products):
            self.matrix(product, f"matrix-sheet-{i}-090", "Zero-Touch Provisioning", "Yes")
        rows = self.rows(include_extended=True)
        self.assertFalse(any(r["name"] == "Zero-Touch Provisioning" for r in rows.values()))
        self.assertFalse(any(
            r["name"] == "Zero-Touch Provisioning"
            for r in self.rows(only_differences=True, include_extended=True).values()
        ))

    def test_workbook_poe_is_visible_with_unknown_standard_placeholder(self):
        self.matrix(self.products[0], "matrix-access-020", "Total PoE Budget", "120 W")
        definition = SpecDefinition.objects.get(code="poe_budget_w")
        ProductSpec.objects.create(product=self.products[0], definition=definition, value_status="unknown")
        self.assertEqual(self.rows()["poe_budget_w"]["values"][0]["display"], "120 W")

    def test_web_page_and_export_use_resolved_values(self):
        self.client.force_login(get_user_model().objects.create_user(username="matrix-viewer"))
        query = {"products": [p.pk for p in self.products]}
        page = self.client.get(reverse("comparison:compare"), query)
        self.assertContains(page, "128 Gbps")
        response = self.client.get(reverse("comparison:export"), query)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        row = next(r for r in workbook["Comparison"].values if r[0] == "Switching Capacity")
        self.assertEqual(row[1:3], ("128 Gbps", "176 Gbps"))
        workbook.close()
