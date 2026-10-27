from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image

from catalog.models import Product, ProductSpec
from imports.template_profiles import get_template_profile, iter_template_profiles


class TPLinkProductDatabaseImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("initialize_catalog", verbosity=0)

    def write_matrix(self, path, sheets):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        workbook.save(path)
        workbook.close()

    def test_mobile_app_label_and_supported_value_survive_reimport(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "apps.xlsx"
            self.write_matrix(path, {"Easy Smart": [
                ["Product Image", "", ""],
                ["Brand", "", "tp-link"],
                ["Model", "", "TL-SG108E"],
                ["Product Name", "", "Easy Smart Switch"],
                ["MANAGEMENT", "Omada App", "Yes"],
            ]})
            for _ in range(2):
                call_command("import_tp_link_product_database", path, sheet="Easy Smart", verbosity=0)
                spec = ProductSpec.objects.get(product__model="TL-SG108E", definition__display_name="Mobile APP")
                self.assertEqual(spec.value_text, "Omada APP")

    def test_imports_switches_into_separate_poe_product_types(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "ProductDB data 2.0.xlsx"
            self.write_matrix(
                path,
                {
                    "Easy Smart": [
                        ["Product Image", "", "", ""],
                        ["Brand", "", "tp-link", "tp-link"],
                        ["Model", "", "TL-SG108PE", "TL-SG108E"],
                        ["Product Name", "", "Easy Smart PoE Switch", "Easy Smart Switch"],
                        ["HARDWARE FEATURES", "Interface", "8× GE RJ45 Ports", "8× GE RJ45 Ports"],
                        ["PoE", "PoE Port Configuration", "Ports 1-4", "—"],
                        ["PoE", "PoE Standard", "802.3af / 802.3at", "—"],
                        ["PoE", "Total PoE Budget", "64 W", "—"],
                        ["PERFORMANCE", "Switching Capacity", "16 Gbps", "16 Gbps"],
                        ["SOFTWARE FEATURES", "VLAN", "Yes", "Yes"],
                        ["SOFTWARE FEATURES", "DNS Security", "DNSSEC", "—"],
                    ],
                    "SW Omada - Campus": [
                        ["Product Image", "", "", ""],
                        ["Brand", "", "Omada", "Omada"],
                        ["Model", "", "S6500-48GP6XF", "S6500-48G6XF"],
                        ["Product Name", "", "Campus L3 Managed PoE+ Switch", "Campus L3 Managed Switch"],
                        ["HARDWARE FEATURES", "Interface", "48× GE RJ45 Ports", "48× GE RJ45 Ports"],
                        ["PoE FEATURES", "PoE Port Configuration", "48× PoE+ Ports", "—"],
                        ["PoE FEATURES", "Total PoE Budget", "1440 W", "—"],
                        ["PERFORMANCE", "Switching Capacity", "216 Gbps", "216 Gbps"],
                        ["SOFTWARE FEATURES", "L3 Features", "OSPF / RIP / VRRP", "OSPF / RIP / VRRP"],
                    ],
                },
            )
            call_command(
                "import_tp_link_product_database",
                path,
                sheet="Easy Smart",
                verbosity=0,
            )
            call_command(
                "import_tp_link_product_database",
                path,
                sheet="SW Omada - Campus",
                verbosity=0,
            )

        expected = {
            "TL-SG108PE": "easy_smart_poe",
            "TL-SG108E": "easy_smart_non_poe",
            "S6500-48GP6XF": "campus_poe",
            "S6500-48G6XF": "campus_non_poe",
        }
        products = {
            product.model: product
            for product in Product.objects.filter(model__in=expected).select_related("product_type")
        }
        self.assertEqual(
            {model: product.product_type.code for model, product in products.items()},
            expected,
        )
        self.assertEqual(products["TL-SG108PE"].product_model.product_name, "Easy Smart PoE Switch")
        self.assertTrue(
            ProductSpec.objects.filter(
                product=products["TL-SG108PE"],
                definition__code="poe_budget_w",
                value_number=64,
            ).exists()
        )
        self.assertFalse(
            ProductSpec.objects.filter(
                product=products["TL-SG108E"],
                definition__code="poe_budget_w",
            ).exists()
        )
        self.assertTrue(
            ProductSpec.objects.filter(
                product=products["TL-SG108PE"],
                definition__code="matrix-easy-smart-011",
                definition__display_name="DNS Security",
                definition__group="SOFTWARE FEATURES",
                value_text="DNSSEC",
            ).exists()
        )

    def test_imports_product_image_from_worksheet_drawing(self):
        with TemporaryDirectory() as directory:
            directory_path = Path(directory)
            workbook_path = directory_path / "ProductDB data 2.0.xlsx"
            image_path = directory_path / "switch.png"
            Image.new("RGB", (8, 8), "white").save(image_path)
            self.write_matrix(
                workbook_path,
                {
                    "Unmanaged": [
                        ["Product Image", "", ""],
                        ["Brand", "", "tp-link"],
                        ["Model", "", "TL-SG105"],
                        ["Product Name", "", "5-Port Desktop Switch"],
                        ["HARDWARE FEATURES", "Interface", "5× GE RJ45 Ports"],
                    ]
                },
            )
            workbook = load_workbook(workbook_path)
            workbook["Unmanaged"].add_image(WorksheetImage(image_path), "C1")
            workbook.save(workbook_path)
            workbook.close()

            with self.settings(MEDIA_ROOT=directory_path / "media"):
                call_command(
                    "import_tp_link_product_database",
                    workbook_path,
                    sheet="Unmanaged",
                    verbosity=0,
                )
                product = Product.objects.get(model="TL-SG105")
                self.assertTrue(product.image.name.startswith("products/"))
                self.assertTrue(Path(product.image.path).is_file())

    def test_derives_per_band_mimo_from_antenna_counts(self):
        with TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "ProductDB data 2.0.xlsx"
            self.write_matrix(
                workbook_path,
                {
                    "WiFi - Ceiling": [
                        ["Product Image", "", ""],
                        ["Brand", "", "Omada"],
                        ["Model", "", "EAP-MIMO"],
                        ["Product Name", "", "Tri-Band Access Point"],
                        [
                            "HARDWARE FEATURES",
                            "Antenna",
                            "2.4 GHz: 2× 4 dBi\n5 GHz: 4× 5 dBi\n6 GHz: 2× 5 dBi",
                        ],
                    ]
                },
            )
            call_command(
                "import_tp_link_product_database",
                workbook_path,
                sheet="WiFi - Ceiling",
                verbosity=0,
            )

        product = Product.objects.get(model="EAP-MIMO")
        specs = {
            spec.definition.code: spec.display_value
            for spec in product.specs.select_related("definition")
        }
        self.assertEqual(specs["mimo_2g"], "2×2 MIMO")
        self.assertEqual(specs["mimo_5g"], "4×4 MIMO")
        self.assertEqual(specs["mimo_6g"], "2×2 MIMO")
        self.assertEqual(specs["total_spatial_streams"], "8")

    def test_dry_run_rolls_back_products(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "ProductDB data 2.0.xlsx"
            self.write_matrix(
                path,
                {
                    "Unmanaged PoE": [
                        ["Product Image", "", ""],
                        ["Brand", "", "tp-link"],
                        ["Model", "", "TL-SG1005P"],
                        ["Product Name", "", "5-Port Switch with 4-Port PoE+"],
                        ["HARDWARE FEATURES", "PoE Port", "Ports 1-4"],
                    ]
                },
            )
            call_command(
                "import_tp_link_product_database",
                path,
                sheet="Unmanaged PoE",
                dry_run=True,
                verbosity=0,
            )
        self.assertFalse(Product.objects.filter(model="TL-SG1005P").exists())

    def test_every_switch_product_type_has_its_own_template_file(self):
        profiles = [
            profile
            for profile in iter_template_profiles()
            if profile.category_slug in {
                "managed-switches",
                "unmanaged-easy-smart-switches",
            }
            and profile.product_type_code != "unknown"
        ]
        self.assertEqual(len(profiles), 14)
        self.assertTrue(all(profile.path.is_file() for profile in profiles))

        poe = get_template_profile("managed-switches", "l2_poe")
        non_poe = get_template_profile("managed-switches", "l2_non_poe")
        poe_codes = {field["code"] for field in poe.fields}
        non_poe_codes = {field["code"] for field in non_poe.fields}
        self.assertIn("poe_budget_w", poe_codes)
        self.assertNotIn("poe_budget_w", non_poe_codes)
