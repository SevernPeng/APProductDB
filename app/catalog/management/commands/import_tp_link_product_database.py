import re
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from catalog.models import (
    Brand,
    Category,
    Product,
    ProductSpec,
    ProductType,
    SpecDefinition,
    normalize_model_key,
)
from catalog.product_types import PRODUCT_TYPE_DEFINITIONS, infer_product_type_code
from catalog.regions import canonical_product_region


SOURCE_VERIFIED_DATE = date(2026, 8, 26)
IGNORED_SHEETS = {"Overview", "Omada Controller"}

SHEET_CONFIG = {
    "Fusion Pro": ("gateway", "integrated_gateway"),
    "Fusion Series": ("gateway", "integrated_gateway"),
    "Gateway - Wired": ("gateway", "wired_router"),
    "Gateway - WiFi": ("gateway", "wireless_router"),
    "Gateway - 4G 5G": ("gateway", "cellular_router"),
    "WiFi - Ceiling": ("access-point", "ceiling"),
    "WiFi - Wall Plate": ("access-point", "wall_plate"),
    "WiFi - Outdoor": ("access-point", "outdoor"),
    "WiFi - Bridge": ("wireless-bridge", ""),
    "WiFi - Desktop": ("access-point", "desktop"),
    "WiFi - GPON": ("access-point", "other"),
    "SW Omada - Campus": ("managed-switches", "switch"),
    "SW Omada - Aggregation": ("managed-switches", "switch"),
    "SW Omada - Access Max": ("managed-switches", "switch"),
    "SW Omada - Access Pro": ("managed-switches", "switch"),
    "SW Omada - Access Plus": ("managed-switches", "switch"),
    "Access Switch": ("managed-switches", "switch"),
    "Agile Switch": ("managed-switches", "switch"),
    "Industrial Switch": ("managed-switches", "switch"),
    "Easy Smart": ("unmanaged-easy-smart-switches", "switch"),
    "Unmanaged": ("unmanaged-easy-smart-switches", "switch"),
    "Unmanaged PoE": ("unmanaged-easy-smart-switches", "switch"),
}

TEXT_SPEC_LABELS = {
    "interface": "ethernet_interfaces",
    "antenna": "antenna_type",
    "wireless security": "wireless_security",
    "frequency": "supported_bands",
    "transmission distance": "wireless_range",
    "poe standard": "poe_standard",
    "standard": "poe_standard",
    "poe port configuration": "poe_configuration",
    "poe port": "poe_configuration",
    "mac address table": "mac_address_table",
    "l2 and l2+ features": "l2_features",
    "l2 features": "l2_features",
    "l3 features": "l3_features",
    "security": "acl_security",
    "access control list": "acl_security",
    "acl": "acl_security",
    "management features": "management_methods",
    "centralized management": "management_methods",
    "installation": "installation",
    "mounting": "installation",
    "operating temperature": "operating_temperature_c",
    "weatherproof enclosure": "ip_rating",
    "surge protection": "lightning_protection_kv",
    "lightning protection": "lightning_protection_kv",
    "power supply": "power_supply",
    "sd-wan": "sd_wan",
    "qos": "qos_support",
    "quality of service": "qos_support",
    "igmp snooping": "igmp_snooping",
    "link aggregation": "link_aggregation",
    "port mirroring": "port_mirroring",
    "cable diagnostics": "cable_test",
    "flow control": "flow_control",
}

BOOLEAN_SPEC_LABELS = {
    "poe auto recovery": "poe_auto_recovery",
    "extend mode": "extend_mode",
    "port isolation": "port_isolation",
    "loop prevention": "loop_prevention",
    "m-lag": "m_lag_support",
    "ptp": "ptp_support",
}


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(value).replace("\u00a0", " ")).strip()


def is_blank_value(value):
    text = clean_text(value)
    return not text or bool(re.fullmatch(r"[—–\-]+", text))


def normalized_label(value):
    return re.sub(r"\s+", " ", clean_text(value)).casefold().rstrip(".")


def extract_number(value, unit=""):
    text = clean_text(value).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        number = Decimal(match.group())
    except InvalidOperation:
        return None
    lowered = text.casefold()
    if unit == "gbps":
        if "tbps" in lowered:
            number *= 1000
        elif "mbps" in lowered:
            number /= 1000
    elif unit == "mpps" and "kpps" in lowered:
        number /= 1000
    elif unit == "mb":
        if "mbit" in lowered or "m-bit" in lowered:
            number /= 8
        elif "kb" in lowered:
            number /= 1024
    elif unit == "bytes":
        if "kb" in lowered:
            number *= 1024
    return number


def parse_boolean(value):
    text = clean_text(value).casefold()
    if is_blank_value(value):
        return None
    if text in {"no", "false", "not supported"}:
        return False
    return True


def parse_poe_ports(value):
    text = clean_text(value)
    if is_blank_value(text):
        return None
    ranges = re.findall(r"ports?\s*(\d+)\s*[-–]\s*(\d+)", text, re.I)
    if ranges:
        return sum(max(int(end) - int(start) + 1, 0) for start, end in ranges)
    counts = re.findall(r"(\d+)\s*(?:×|x|-)?\s*poe", text, re.I)
    if counts:
        return sum(int(item) for item in counts)
    count = re.search(r"(\d+)[- ]*port", text, re.I)
    return int(count.group(1)) if count else None


def parse_wireless_rates(value):
    text = clean_text(value).replace(",", "")
    rates = {}
    for band, code in (("2.4", "rate_2g_mbps"), ("5", "rate_5g_mbps"), ("6", "rate_6g_mbps")):
        match = re.search(
            rf"{re.escape(band)}\s*GHz[^\d]{{0,30}}(\d+(?:\.\d+)?)\s*Mbps",
            text,
            re.I,
        )
        if match:
            rates[code] = Decimal(match.group(1))
    return rates


def parse_antenna_mimo(value):
    """Infer per-radio MIMO from the antenna count published for each band."""
    text = clean_text(value)
    mimo = {}
    for band, code in (("2.4", "mimo_2g"), ("5", "mimo_5g"), ("6", "mimo_6g")):
        match = re.search(
            rf"{re.escape(band)}\s*GHz\s*:?\s*(?:[•·]\s*)?"
            r"(\d+)\s*[x×]\s*\d+(?:\.\d+)?\s*dBi",
            text,
            re.I,
        )
        if match:
            streams = int(match.group(1))
            mimo[code] = f"{streams}×{streams} MIMO"
    if not mimo:
        # The legacy 2.4 GHz N-series rows omit the band label because they
        # contain only one radio.
        single_radio = re.fullmatch(
            r"\s*(\d+)\s*[x×]\s*\d+(?:\.\d+)?\s*dBi(?:\s+Omni)?\s*",
            text,
            re.I,
        )
        if single_radio:
            streams = int(single_radio.group(1))
            mimo["mimo_2g"] = f"{streams}×{streams} MIMO"
    if mimo:
        mimo["total_spatial_streams"] = sum(
            int(re.search(r"\d+", value).group())
            for value in mimo.values()
        )
    return mimo


def has_poe_value(rows, column):
    for _group, label, values in rows:
        if normalized_label(label) not in {"poe port", "poe port configuration"}:
            continue
        return not is_blank_value(values[column])
    return False


def switch_family(sheet_name, model, product_name):
    if sheet_name == "SW Omada - Campus":
        return "campus"
    if sheet_name == "Agile Switch":
        return "agile"
    if sheet_name == "Industrial Switch":
        return "industrial"
    if sheet_name == "Easy Smart":
        return "easy_smart"
    if sheet_name in {"Unmanaged", "Unmanaged PoE"}:
        return "unmanaged"
    if re.search(r"(?:\blite\s+)?\bl3\b|layer\s*3", product_name, re.I):
        return "l3"
    if re.match(r"^(?:S[67]|SX6|SG6|SG5)", model, re.I):
        return "l3"
    return "l2"


def product_type_for(sheet_name, category_slug, configured_type, model, product_name, has_poe):
    if configured_type != "switch":
        return configured_type
    family = switch_family(sheet_name, model, product_name)
    if category_slug == "unmanaged-easy-smart-switches":
        return f"{family}_{'poe' if has_poe else 'non_poe'}"
    return infer_product_type_code(
        category_slug,
        model,
        has_poe=has_poe,
        switch_family=family,
        product_name=product_name,
    )


def source_rows(sheet):
    rows = []
    for row in sheet.iter_rows(values_only=True):
        group = clean_text(row[0] if len(row) > 0 else "")
        label = clean_text(row[1] if len(row) > 1 else "")
        rows.append((group, label, row))
    return rows


def matrix_spec_rows(sheet):
    """Yield the workbook's visible group, label, value rows in source order."""
    current_group = "Other Specifications"
    for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=5):
        group = clean_text(row[0] if len(row) > 0 else "")
        label = clean_text(row[1] if len(row) > 1 else "")
        if group:
            current_group = group
        if label:
            yield row_number, current_group, label, row


def matrix_definition_code(sheet_name, row_number):
    sheet_key = slugify(sheet_name)[:70] or "sheet"
    return f"matrix-{sheet_key}-{row_number:03d}"


def _resolve_xlsx_target(source_path, target):
    return posixpath.normpath(
        posixpath.join(posixpath.dirname(source_path), target)
    ).lstrip("/")


def _relationships(archive, source_path):
    rels_path = posixpath.join(
        posixpath.dirname(source_path),
        "_rels",
        posixpath.basename(source_path) + ".rels",
    )
    if rels_path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rels_path))
    return {
        item.attrib["Id"]: {
            "type": item.attrib.get("Type", ""),
            "target": _resolve_xlsx_target(source_path, item.attrib["Target"]),
        }
        for item in root
    }


def workbook_product_images(workbook_path):
    """Extract both drawing images and Excel rich-data in-cell images."""
    images = {}
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    drawing_main_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"

    with zipfile.ZipFile(workbook_path) as archive:
        workbook_path_in_zip = "xl/workbook.xml"
        workbook_root = ET.fromstring(archive.read(workbook_path_in_zip))
        workbook_rels = _relationships(archive, workbook_path_in_zip)
        sheet_paths = {
            sheet.attrib["name"]: workbook_rels[
                sheet.attrib[f"{{{office_rel_ns}}}id"]
            ]["target"]
            for sheet in workbook_root.findall(f".//{{{main_ns}}}sheet")
        }

        rich_images = {}
        metadata_path = "xl/metadata.xml"
        rich_values_path = "xl/richData/rdrichvalue.xml"
        rich_value_rels_path = "xl/richData/richValueRel.xml"
        if all(
            path in archive.namelist()
            for path in (metadata_path, rich_values_path, rich_value_rels_path)
        ):
            metadata_root = ET.fromstring(archive.read(metadata_path))
            metadata_indexes = [
                int(next(iter(block.iterfind(".//{*}rvb"))).attrib["i"])
                for block in metadata_root.findall(".//{*}futureMetadata/{*}bk")
            ]
            rich_values_root = ET.fromstring(archive.read(rich_values_path))
            rich_value_rel_indexes = [
                int(item.find("{*}v").text)
                for item in rich_values_root.findall("{*}rv")
            ]
            rich_rel_root = ET.fromstring(archive.read(rich_value_rels_path))
            rich_rel_ids = [
                item.attrib[f"{{{office_rel_ns}}}id"] for item in rich_rel_root
            ]
            rich_rels = _relationships(archive, rich_value_rels_path)
            for metadata_position, rich_value_index in enumerate(metadata_indexes, start=1):
                rel_position = rich_value_rel_indexes[rich_value_index]
                image_path = rich_rels[rich_rel_ids[rel_position]]["target"]
                rich_images[metadata_position] = (
                    Path(image_path).suffix.lower() or ".png",
                    archive.read(image_path),
                )

        for sheet_name, sheet_path in sheet_paths.items():
            sheet_root = ET.fromstring(archive.read(sheet_path))
            for cell in sheet_root.findall(f".//{{{main_ns}}}row[@r='1']/{{{main_ns}}}c"):
                metadata_position = cell.attrib.get("vm")
                if not metadata_position or int(metadata_position) not in rich_images:
                    continue
                column_letters = re.match(r"[A-Z]+", cell.attrib["r"]).group()
                column = 0
                for letter in column_letters:
                    column = column * 26 + ord(letter) - 64
                images[(sheet_name, column - 1)] = rich_images[int(metadata_position)]

            sheet_rels = _relationships(archive, sheet_path)
            for relation in sheet_rels.values():
                if not relation["type"].endswith("/drawing"):
                    continue
                drawing_path = relation["target"]
                drawing_root = ET.fromstring(archive.read(drawing_path))
                drawing_rels = _relationships(archive, drawing_path)
                for anchor in list(drawing_root):
                    start = anchor.find(f"{{{drawing_ns}}}from")
                    blip = anchor.find(f".//{{{drawing_main_ns}}}blip")
                    if start is None or blip is None:
                        continue
                    row = int(start.find(f"{{{drawing_ns}}}row").text)
                    column = int(start.find(f"{{{drawing_ns}}}col").text)
                    if row != 0:
                        continue
                    rel_id = blip.attrib[f"{{{office_rel_ns}}}embed"]
                    image_path = drawing_rels[rel_id]["target"]
                    images[(sheet_name, column)] = (
                        Path(image_path).suffix.lower() or ".png",
                        archive.read(image_path),
                    )
    return images


def find_row(rows, first_cell):
    target = first_cell.casefold()
    for _group, _label, values in rows:
        if clean_text(values[0] if values else "").casefold() == target:
            return values
    return None


def spec_values_for(category_slug, rows, column):
    values = {}
    for group, label, row in rows:
        value = row[column] if column < len(row) else None
        if is_blank_value(value):
            continue
        key = normalized_label(label)
        source_value = clean_text(value)
        if key in {"fan", "fan quantity"}:
            if category_slug == "unmanaged-easy-smart-switches":
                values["fanless"] = bool(
                    re.search(r"fanless|no fan", source_value, re.I)
                )
            else:
                values["fan_design"] = source_value
            continue
        if key == "signal rate":
            values.update(parse_wireless_rates(source_value))
            continue
        if key == "antenna" and category_slug == "access-point":
            values.update(parse_antenna_mimo(source_value))
        if key in {"concurrent clients", "client capacity"}:
            number = extract_number(source_value)
            if number is not None:
                values["max_clients"] = number
            continue
        numeric = {
            "switching capacity": ("switching_capacity_gbps", "gbps"),
            "packet forwarding rate": ("packet_forwarding_rate_mpps", "mpps"),
            "packet buffer": ("packet_buffer_mb", "mb"),
            "packet buffer memory": ("packet_buffer_mb", "mb"),
            "jumbo frame": ("jumbo_frame_bytes", "bytes"),
            "total poe budget": ("poe_budget_w", "w"),
            "poe budget": ("router_poe_budget_w", "w"),
            "max poe output per port": ("max_poe_per_port_w", "w"),
            "max. poe output per port": ("max_poe_per_port_w", "w"),
            "max power consumption": ("max_power_consumption_w", "w"),
            "max. power consumption": ("max_power_consumption_w", "w"),
            "concurrent sessions": ("concurrent_sessions", ""),
            "concurrent session": ("concurrent_sessions", ""),
            "nat throughput (static ip)": ("nat_throughput_mbps", ""),
            "nat (static ip)": ("nat_throughput_mbps", ""),
            "ipsec vpn throughput": ("vpn_throughput_mbps", ""),
            "vpn throughput": ("vpn_throughput_mbps", ""),
        }.get(key)
        if numeric:
            number = extract_number(source_value, numeric[1])
            if number is not None:
                values[numeric[0]] = max(values.get(numeric[0], number), number)
            continue
        if key in {"poe port", "poe port configuration"}:
            count = parse_poe_ports(source_value)
            if count is not None:
                code = "router_poe_ports" if category_slug == "gateway" else "poe_ports"
                values[code] = Decimal(count)
        text_code = TEXT_SPEC_LABELS.get(key)
        if text_code:
            if key == "power supply" and category_slug in {"access-point", "wireless-bridge"}:
                text_code = "power_method"
            elif key == "centralized management" and category_slug == "gateway":
                text_code = "controller_management"
            elif key == "centralized management" and category_slug in {"access-point", "wireless-bridge"}:
                text_code = "centralized_management"
            elif text_code == "poe_standard" and category_slug == "gateway":
                continue
            previous = values.get(text_code)
            values[text_code] = f"{previous}\n{source_value}" if previous else source_value
            continue
        boolean_code = BOOLEAN_SPEC_LABELS.get(key)
        if boolean_code:
            boolean = parse_boolean(source_value)
            if boolean is not None:
                values[boolean_code] = boolean
    return values


def typed_spec_value(definition, value):
    result = {"value_text": "", "value_number": None, "value_boolean": None}
    if definition.data_type in {SpecDefinition.DataType.INTEGER, SpecDefinition.DataType.DECIMAL}:
        quantum = Decimal("1") if definition.data_type == SpecDefinition.DataType.INTEGER else Decimal("0.001")
        result["value_number"] = Decimal(value).quantize(quantum, rounding=ROUND_HALF_UP)
    elif definition.data_type == SpecDefinition.DataType.BOOLEAN:
        result["value_boolean"] = bool(value)
    else:
        result["value_text"] = clean_text(value)
    return result


class Command(BaseCommand):
    help = "Import the TP-Link/Omada ProductDB 2.0 matrix workbook into the catalog."

    def add_arguments(self, parser):
        parser.add_argument("workbook", type=Path)
        parser.add_argument("--sheet", action="append", dest="sheets")
        parser.add_argument("--dry-run", action="store_true")

    @transaction.atomic
    def handle(self, *args, **options):
        workbook_path = options["workbook"].resolve()
        if not workbook_path.is_file():
            raise CommandError(f"Workbook not found: {workbook_path}")
        product_images = workbook_product_images(workbook_path)
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        selected = options["sheets"] or [name for name in workbook.sheetnames if name not in IGNORED_SHEETS]
        if isinstance(selected, str):
            selected = [selected]
        unsupported = set(selected) - set(SHEET_CONFIG)
        missing = set(selected) - set(workbook.sheetnames)
        if unsupported:
            workbook.close()
            raise CommandError(f"Unsupported sheets: {', '.join(sorted(unsupported))}")
        if missing:
            workbook.close()
            raise CommandError(f"Missing sheets: {', '.join(sorted(missing))}")

        categories = {item.slug: item for item in Category.objects.all()}
        for slug in {SHEET_CONFIG[name][0] for name in selected}:
            if slug not in categories:
                workbook.close()
                raise CommandError(f"Catalog category is not initialized: {slug}")
        for category_slug, definitions in PRODUCT_TYPE_DEFINITIONS.items():
            category = categories[category_slug]
            for order, (code, name, description) in enumerate(definitions, start=1):
                ProductType.objects.update_or_create(
                    category=category,
                    code=code,
                    defaults={"name": name, "description": description, "display_order": order * 10, "active": True},
                )

        product_types = {
            (item.category.slug, item.code): item
            for item in ProductType.objects.select_related("category").filter(active=True)
        }
        definitions = {item.code: item for item in SpecDefinition.objects.filter(active=True)}
        brand, _ = Brand.objects.update_or_create(
            slug="tp-link",
            defaults={"name": "TP-Link", "is_own_brand": True, "active": True},
        )
        counts = Counter()
        source_note = f"Imported from {workbook_path.name}"

        for sheet_name in selected:
            category_slug, configured_type = SHEET_CONFIG[sheet_name]
            category = categories[category_slug]
            sheet = workbook[sheet_name]
            rows = source_rows(sheet)
            raw_rows = list(matrix_spec_rows(sheet))
            model_row = find_row(rows, "Model")
            product_name_row = find_row(rows, "Product Name")
            if model_row is None or product_name_row is None:
                workbook.close()
                raise CommandError(f"{sheet_name} is missing the Model or Product Name row")
            for column in range(2, len(model_row)):
                model = clean_text(model_row[column])
                if not model:
                    continue
                product_name = clean_text(product_name_row[column] if column < len(product_name_row) else "")
                has_poe = has_poe_value(rows, column)
                type_code = product_type_for(
                    sheet_name,
                    category_slug,
                    configured_type,
                    model,
                    product_name,
                    has_poe,
                )
                product_type = product_types.get((category_slug, type_code)) if type_code else None
                if type_code and product_type is None:
                    raise CommandError(f"Missing product type {category_slug}/{type_code} for {model}")
                region = canonical_product_region(category_slug) or "UN"
                product = (
                    Product.objects.filter(
                        brand=brand,
                        model_key=normalize_model_key(model),
                        region=region,
                    )
                    .order_by("pk")
                    .first()
                )
                created = product is None
                if created:
                    product = Product.objects.create(
                        brand=brand,
                        category=category,
                        product_type=product_type,
                        model=model,
                        region=region,
                        ap_type=type_code if category_slug == "access-point" and type_code != "other" else Product.APType.OTHER if category_slug == "access-point" else "",
                        lifecycle_status=Product.LifecycleStatus.ACTIVE,
                        notes=source_note,
                        is_published=True,
                    )
                else:
                    changed = []
                    if product.category_id == category.pk and product.product_type_id != getattr(product_type, "pk", None):
                        product.product_type = product_type
                        changed.append("product_type")
                    if not product.is_published:
                        product.is_published = True
                        changed.append("is_published")
                    if source_note not in (product.notes or ""):
                        product.notes = "\n".join(filter(None, (product.notes, source_note)))
                        changed.append("notes")
                    if changed:
                        product.save(update_fields=tuple(changed) + ("updated_at",))
                counts["products_created" if created else "products_updated"] += 1
                if category_slug in {
                    "managed-switches",
                    "unmanaged-easy-smart-switches",
                }:
                    counts[f"switch_type_{type_code}"] += 1

                if product.product_model_id:
                    product_model = product.product_model
                    model_changed = []
                    if product_name and product_model.product_name != product_name:
                        product_model.product_name = product_name
                        model_changed.append("product_name")
                    series = sheet_name.replace("SW Omada - ", "").replace("WiFi - ", "")
                    if product_model.series != series:
                        product_model.series = series
                        model_changed.append("series")
                    if model_changed:
                        product_model.save(update_fields=tuple(model_changed) + ("updated_at",))

                image_data = product_images.get((sheet_name, column))
                if image_data and not options["dry_run"]:
                    suffix, content = image_data
                    current_content = b""
                    if product.image:
                        try:
                            with product.image.open("rb") as current_image:
                                current_content = current_image.read()
                        except OSError:
                            current_content = b""
                    if current_content != content:
                        product.image.save(
                            f"{slugify(model) or 'product'}{suffix}",
                            ContentFile(content),
                            save=True,
                        )
                        counts["images_saved"] += 1
                    else:
                        counts["images_unchanged"] += 1

                imported_matrix_codes = []
                for row_number, group, label, raw_row in raw_rows:
                    value = raw_row[column] if column < len(raw_row) else None
                    if is_blank_value(value):
                        continue
                    if " ".join(label.split()).casefold() in {"omada app", "mobile app"}:
                        label = "Mobile APP"
                        if clean_text(value).casefold() == "yes":
                            value = "Omada APP"
                    code = matrix_definition_code(sheet_name, row_number)
                    definition, definition_created = SpecDefinition.objects.update_or_create(
                        code=code,
                        defaults={
                            "display_name": label,
                            "group": group,
                            "category": category,
                            "data_type": SpecDefinition.DataType.TEXT,
                            "is_core": False,
                            "display_order": row_number * 10,
                            "description": f"{sheet_name}: {group} / {label}",
                            "active": True,
                        },
                    )
                    definitions[code] = definition
                    imported_matrix_codes.append(code)
                    _matrix_spec, matrix_spec_created = ProductSpec.objects.update_or_create(
                        product=product,
                        definition=definition,
                        defaults={
                            "value_text": clean_text(value),
                            "value_number": None,
                            "value_boolean": None,
                            "value_status": ProductSpec.ValueStatus.PUBLISHED,
                            "unit": "",
                            "raw_value": clean_text(value),
                            "source_note": (
                                f"{source_note}; {sheet_name}!"
                                f"{get_column_letter(column + 1)}{row_number}"
                            ),
                            "verified_date": SOURCE_VERIFIED_DATE,
                        },
                    )
                    counts[
                        "matrix_specs_created" if matrix_spec_created else "matrix_specs_updated"
                    ] += 1
                    counts[
                        "matrix_definitions_created" if definition_created else "matrix_definitions_updated"
                    ] += 1

                matrix_prefix = f"matrix-{slugify(sheet_name)[:70] or 'sheet'}-"
                stale_specs = product.specs.filter(
                    definition__code__startswith=matrix_prefix
                ).exclude(definition__code__in=imported_matrix_codes)
                deleted_specs, _ = stale_specs.delete()
                counts["matrix_specs_deleted"] += deleted_specs

                spec_values = spec_values_for(category_slug, rows, column)
                if category_slug == "unmanaged-easy-smart-switches":
                    spec_values["management_type"] = "Easy Smart" if type_code.startswith("easy_smart") else "Unmanaged"
                for code, value in spec_values.items():
                    definition = definitions.get(code)
                    if definition is None:
                        continue
                    typed = typed_spec_value(definition, value)
                    _spec, spec_created = ProductSpec.objects.update_or_create(
                        product=product,
                        definition=definition,
                        defaults={
                            **typed,
                            "value_status": ProductSpec.ValueStatus.PUBLISHED,
                            "unit": definition.unit,
                            "raw_value": clean_text(value),
                            "source_note": f"{source_note}; {sheet_name}!{sheet.cell(row=3, column=column + 1).coordinate}",
                            "verified_date": SOURCE_VERIFIED_DATE,
                        },
                    )
                    counts["specs_created" if spec_created else "specs_updated"] += 1

        workbook.close()
        if options["dry_run"]:
            transaction.set_rollback(True)
        mode = "DRY RUN" if options["dry_run"] else "IMPORTED"
        self.stdout.write(self.style.SUCCESS(f"{mode}: {dict(sorted(counts.items()))}"))
